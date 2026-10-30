#!/usr/bin/env python3
"""Export the complete, editable ten-lab Markdown manuscript to A4 landscape Word."""
from __future__ import annotations

import re
import subprocess
import tempfile
from pathlib import Path

import fitz
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "exports/print/AIClassDocuments-TongHop-10-Lab-BaiTap.md"
WORD = ROOT / "exports/print/source/AIClassDocuments-TongHop-10-Lab-BaiTap.docx"
PARTS = [
    "Đề bài tiếng Anh",
    "Đề bài tiếng Việt",
    "Code Python có chú thích",
    "Giải thích nội dung và cách giải",
    "Giải thích các định nghĩa",
]

def reference_doc(path: Path) -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(29.7), Cm(21)
    sec.left_margin, sec.right_margin = Cm(1.45), Cm(1.45)
    sec.top_margin, sec.bottom_margin = Cm(1.25), Cm(1.25)
    sec.footer_distance = Cm(.6)
    style = doc.styles["Normal"]
    style.font.name, style.font.size = "DejaVu Sans", Pt(8.8)
    style.paragraph_format.space_after = Pt(4)
    style.paragraph_format.line_spacing = 1.1
    for i, size in [(1, 17), (2, 12), (3, 10), (4, 9)]:
        style = doc.styles[f"Heading {i}"]
        style.font.name, style.font.size = "DejaVu Sans", Pt(size)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(8)
        style.paragraph_format.space_after = Pt(4)
    doc.styles["Heading 1"].paragraph_format.page_break_before = True
    source_code = (doc.styles["Source Code"] if "Source Code" in doc.styles
                   else doc.styles.add_style("Source Code", WD_STYLE_TYPE.PARAGRAPH))
    source_code.font.name = "DejaVu Sans Mono"
    source_code.font.size = Pt(7)
    source_code.paragraph_format.space_before = Pt(0)
    source_code.paragraph_format.space_after = Pt(0)
    source_code.paragraph_format.line_spacing = 1
    footer = sec.footer.paragraphs[0]
    footer.add_run("AIClassDocuments  |  ")
    footer.alignment = 2
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    footer._p.append(fld)
    doc.save(path)

def validate(docx_path: Path, source: str) -> None:
    word = Document(docx_path)
    paragraphs = word.paragraphs
    chapter_titles = [p.text for p in paragraphs if p.style.name == "Heading 1"]
    if len(chapter_titles) != 11:
        raise ValueError(f"Expected cover and ten chapters, got {len(chapter_titles)}")
    for section in PARTS:
        count = sum(section in p.text and p.style.name.startswith("Heading")
                    for p in paragraphs)
        if count != 10:
            raise ValueError(f"Missing or duplicate Word section {section}: {count}")
    whole_text = "\n".join(p.text for p in paragraphs)
    if len(whole_text) < 100000 or len(source) < 150000:
        raise ValueError("Word manuscript is unexpectedly short")
    if "def analyzeClass()" not in whole_text or "def task04_train(" not in whole_text:
        raise ValueError("Code sections were lost in Word conversion")
    if word.sections[0].page_width <= word.sections[0].page_height:
        raise ValueError("Word page size must be A4 landscape")

def render_qa(word: Path, directory: Path) -> int:
    subprocess.run(
        ["libreoffice", "-env:UserInstallation=file:///tmp/ai_labs_docx_qa",
         "--headless", "--convert-to", "pdf", "--outdir", str(directory), str(word)],
        check=True, capture_output=True, text=True, timeout=240)
    pdf_name = directory / (word.stem + ".pdf")
    with fitz.open(pdf_name) as pdf:
        if pdf.page_count < 30:
            raise ValueError("Word rendered too few pages")
        for i, page in enumerate(pdf):
            if abs(page.rect.width - 841.89) > 2 or abs(page.rect.height - 595.28) > 2:
                raise ValueError(f"Word page {i+1} is not A4 landscape")
            for block in page.get_text("blocks"):
                x0, _, x1, _ = block[:4]
                if x0 < -1 or x1 > page.rect.width + 1:
                    raise ValueError(f"Clipped text on page {i+1}")
        text = "\n".join(p.get_text("text") for p in pdf)
        for phrase in PARTS:
            if phrase not in text:
                raise ValueError(f"Word render omitted {phrase}")
        return pdf.page_count

def main() -> None:
    markdown = SOURCE.read_text(encoding="utf-8")
    for part in PARTS:
        pattern = "^## [1-5]\\. " + re.escape(part) + "$"
        if len(re.findall(pattern, markdown, re.MULTILINE)) != 10:
            raise ValueError(f"Markdown does not have ten sections named {part}")
    WORD.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="lab_word_") as temp:
        temp_dir = Path(temp)
        ref = temp_dir / "reference.docx"
        reference_doc(ref)
        subprocess.run(
            ["pandoc", str(SOURCE), "--from=gfm", "--to=docx", "--standalone",
             "--reference-doc", str(ref), "--output", str(WORD)],
            check=True, capture_output=True, text=True, timeout=160)
        validate(WORD, markdown)
        pages = render_qa(WORD, temp_dir)
    print(f"WORD_SOURCE_OK chapters=10 sections=50 pages={pages} "
          f"bytes={WORD.stat().st_size} output={WORD.relative_to(ROOT)}")

if __name__ == "__main__":
    main()
